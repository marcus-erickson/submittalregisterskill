#!/usr/bin/env python3
"""Checks the submittal-intake script on the health sciences example register.

Run from the repository root:  python3 tests/test_intake.py
"""
import csv
import json
import os
import subprocess
import sys
import tempfile

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
BUILD = os.path.join(ROOT, "submittal-register", "scripts", "build_register.py")
INTAKE = os.path.join(ROOT, "submittal-intake", "scripts", "intake.py")
EX = os.path.join(ROOT, "submittal-register", "examples", "health-sciences-drainage")
TMP = tempfile.mkdtemp(prefix="intaketest_")
REG = os.path.join(TMP, "register.json")
LOG = os.path.join(TMP, "intake-log.csv")
failures = []


def check(name, ok, detail=""):
    print(("PASS  " if ok else "FAIL  ") + name + (f"  {detail}" if detail and not ok else ""))
    if not ok:
        failures.append(name)


def run(*args):
    r = subprocess.run([sys.executable, *args], capture_output=True, text=True)
    return r


def item(sub_no):
    return next(i for i in json.load(open(REG))["items"] if i["sub_no"] == sub_no)


def run_events(events, tag, dry=False):
    ev, mt = os.path.join(TMP, f"{tag}-events.json"), os.path.join(TMP, f"{tag}-matches.json")
    json.dump({"events": events}, open(ev, "w"))
    m = run(INTAKE, "match", REG, ev, "-o", mt)
    matches = json.load(open(mt))
    args = [INTAKE, "apply", REG, mt, "--log", LOG] + (["--dry-run"] if dry else [])
    return matches["events"], run(*args)


run(BUILD, EX + "-reviewed.csv", os.path.join(TMP, "r.xlsx"), "--rules", EX + "-rules.json", "--save-json", REG)
check("register built", os.path.exists(REG))


def E(i, kind, day, **kw):
    return {"source_id": f"m{i}", "kind": kind, "date": day, "source_name": f"test-{i}.pdf", **kw}


# 1. matching verdicts, no changes yet
events = [
    E(1, "received", "2026-10-01", section="31 20 00", title="Material test reports", types="TEST", reviewer="Architect"),
    E(2, "received", "2026-10-01", section="033053", title="Concrete design mixtures", types="TEST", reviewer="Architect"),
    E(3, "received", "2026-10-01", section="32 13 73", title="joint sealant data", types="PD"),
    E(4, "received", "2026-10-01", title="Engineering survey of condition of the building"),
    E(5, "received", "2026-10-01", section="09 91 23", title="Interior painting product data", types="PD"),
    E(6, "received", "2026-10-01", sub_no="32 12 16-1", reviewer="Architect"),
]
matched, _ = run_events(events, "verdicts", dry=True)
v = {e["source_id"]: e for e in matched}
check("single line in the section matches confidently", v["m1"]["verdict"] == "confident" and v["m1"]["chosen"] == "31 20 00-1")
check("compact section number 033053 normalizes and picks the TEST line",
      v["m2"]["verdict"] == "confident" and v["m2"]["chosen"] == "03 30 53-2", str(v["m2"]["candidates"]))
check("two similar lines in one section are ambiguous, not guessed", v["m3"]["verdict"] == "ambiguous" and not v["m3"]["chosen"],
      str(v["m3"]["candidates"]))
check("same title in two sections, no section given: ambiguous", v["m4"]["verdict"] == "ambiguous" and len(v["m4"]["candidates"]) >= 2)
check("section not in the register: no match, nothing chosen", v["m5"]["verdict"] == "none" and not v["m5"]["chosen"])
check("register number in the email is used", v["m6"]["verdict"] == "confident" and v["m6"]["chosen"] == "32 12 16-1")
check("dry run changed nothing", item("31 20 00-1")["status"] == "Not Started" and not os.path.exists(LOG))

# 1b. the person picked a register, but the document may belong to another project
proj = [
    E(41, "received", "2026-10-01", section="31 20 00", title="Material test reports", types="TEST", project_hint="Westside Elementary Renovation"),
    E(42, "received", "2026-10-01", section="31 20 00", title="Material test reports", types="TEST", project_hint="Northeast Texas Community College, Health Sciences Lab"),
    E(43, "received", "2026-10-01", section="31 20 00", title="Material test reports", types="TEST", project_hint="24-117"),
]
matched, _ = run_events(proj, "project", dry=True)
v2 = {e["source_id"]: e for e in matched}
check("document naming a different project is held back, nothing chosen",
      v2["m41"]["verdict"] == "wrong_project" and not v2["m41"]["chosen"], v2["m41"]["reason"])
check("document naming this project passes", v2["m42"]["verdict"] == "confident")
check("a bare project number is not compared when the register has none", v2["m43"]["verdict"] == "confident")
withnum = json.load(open(REG))
withnum["rules"]["project_number"] = "2025-203"
numreg = os.path.join(TMP, "numbered.json")
json.dump(withnum, open(numreg, "w"))
ev_n = os.path.join(TMP, "num-events.json")
json.dump({"events": [proj[2]]}, open(ev_n, "w"))
run(INTAKE, "match", numreg, ev_n, "-o", os.path.join(TMP, "num-matches.json"))
check("a bare project number that differs from the register's number is held back",
      json.load(open(os.path.join(TMP, "num-matches.json")))["events"][0]["verdict"] == "wrong_project")

# 1c. file ids are stable and content based
f1, f2 = os.path.join(TMP, "a.pdf"), os.path.join(TMP, "b.pdf")
open(f1, "wb").write(b"transmittal one")
open(f2, "wb").write(b"transmittal two")
id1, id1b, id2 = (run(INTAKE, "id", f).stdout.strip() for f in (f1, f1, f2))
check("file id is stable for the same file and differs for different content",
      len(id1) == 16 and id1 == id1b and id1 != id2)

# 2. apply the confident ones: submitted, then returned approved as noted
matched, r = run_events([e for e in events if e["source_id"] in ("m1", "m2", "m6")], "apply1")
check("received sets status, date and ball in court",
      item("31 20 00-1")["status"] == "Submitted" and item("31 20 00-1")["actual_submit"] == "2026-10-01"
      and item("31 20 00-1")["ball_in_court"] == "Architect", r.stdout)
ret = [E(11, "returned", "2026-10-15", sub_no="31 20 00-1", action="approved_as_noted", submitter="Contractor", comment="Furnish test data per note 2")]
run_events(ret, "return1")
i = item("31 20 00-1")
check("returned approved as noted: status, approval date, ball cleared, note kept",
      i["status"] == "Approved as Noted" and i["approval_date"] == "2026-10-15" and i["ball_in_court"] == ""
      and "note 2" in i["notes"], str(i))

# 3. resubmittal cycle on 03 30 53-1, listed out of date order on purpose
cycle = [
    E(23, "received", "2026-10-20", sub_no="03 30 53-1", reviewer="Architect"),
    E(21, "received", "2026-10-01", sub_no="03 30 53-1", reviewer="Architect"),
    E(22, "returned", "2026-10-10", sub_no="03 30 53-1", action="revise_resubmit", submitter="Subcontractor"),
]
run_events(cycle, "cycle")
i = item("03 30 53-1")
check("events apply oldest first, so the resubmittal raises Rev and reopens the line",
      i["rev"] == 1 and i["status"] == "Submitted" and i["actual_submit"] == "2026-10-20" and i["approval_date"] == "", str(i))

# 4. a repeat run of the same emails changes nothing
before = open(REG).read()
_, r = run_events(cycle, "cycle-again")
check("running the same emails again is a no-op", open(REG).read() == before and "skipped 3" in r.stdout, r.stdout)

# 5. conflicts are reported and left alone
before = open(REG).read()
bad = [
    E(31, "returned", "2026-10-12", sub_no="01 31 00-1", action="approved"),  # never submitted
    E(32, "received", "2026-10-25", sub_no="31 20 00-1"),  # already Approved as Noted
    E(33, "returned", "2026-10-25", sub_no="03 30 53-1", action="sort_of_ok"),  # unknown action
]
_, r = run_events(bad, "conflicts")
check("out of order, already-closed and unknown-action events are not applied",
      open(REG).read() == before and "conflicts 2" in r.stdout and "unresolved 1" in r.stdout, r.stdout)
okc = [E(34, "returned", "2026-10-12", sub_no="01 31 00-1", action="approved", confirm=True)]
run_events(okc, "confirmed")
check("a person-confirmed out-of-order change is applied", item("01 31 00-1")["status"] == "Approved")

# 6. unresolved events are not applied and a backup is kept
check("backup of the register before a write exists", os.path.exists(REG + ".bak"))
rows = list(csv.DictReader(open(LOG, encoding="utf-8")))
check("log records applied, conflict and unresolved results",
      {"applied", "conflict", "unresolved", "skipped"} <= {r["result"] for r in rows}, str({r["result"] for r in rows}))
check("nothing was written for the ambiguous and no-match events",
      item("32 13 73-1")["status"] == "Not Started" and item("32 13 73-2")["status"] == "Not Started")

# 7. the updated register still builds a workbook with the new status in it
x = os.path.join(TMP, "after.xlsx")
r = run(BUILD, REG, x)
import openpyxl  # noqa: E402
ws = openpyxl.load_workbook(x)["Register"]
statuses = [c.value for c in ws[4]]
col = statuses.index("Status") + 1
vals = [ws.cell(row=n, column=col).value for n in range(5, ws.max_row + 1)]
check("rebuilt workbook shows the updated statuses", "Approved as Noted" in vals and "Submitted" in vals and r.returncode == 0)

print()
if failures:
    print(f"{len(failures)} check(s) failed")
    sys.exit(1)
print("All checks passed")
