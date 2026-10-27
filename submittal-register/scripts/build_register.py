#!/usr/bin/env python3
"""Build the submittal register workbook from a reviewed CSV.

Usage:
  build_register.py reviewed.csv out.xlsx [--rules rules.json] [--project NAME]
                    [--review-weeks N] [--prep-weeks N] [--numbering section|sequential|compact]
                    [--install-dates dates.csv]

Input CSV columns (all optional except section and description or requirement):
  sub_no, section, section_title, description, requirement, types, article, page,
  origin, source, lead_weeks, review_weeks, notes, needs_review, review_note

The raw output of extract.py works as is. Claude's reviewed file adds sub_no,
lead_weeks, notes, needs_review and review_note.

Install dates file (optional): section,install_date,lead_weeks  (dates as YYYY-MM-DD)
"""
import argparse
import csv
import json
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
STATUSES = ["Not Started", "In Prep", "Submitted", "Approved", "Approved as Noted",
            "Revise/Resubmit", "Rejected"]
TYPE_LEGEND = [
    ("PD", "Product data"), ("SD", "Shop drawings"), ("WD", "Wiring diagrams"),
    ("SAM", "Samples"), ("CALC", "Calculations"), ("CERT", "Certificates, qualifications"),
    ("TEST", "Test and inspection reports"), ("WARR", "Warranties"),
    ("OM", "Operation and maintenance data"), ("MOCKUP", "Mock-ups"),
    ("RECORD", "Record drawings and as-builts"), ("SCHED", "Schedules"),
    ("REPORT", "Surveys, statements, other reports"), ("REF", "Cross-reference to another section"),
    ("OTHER", "Not classified"),
]

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def truthy(v):
    return str(v or "").strip().lower() in ("y", "yes", "true", "1", "x")


def shorten(text, n=240):
    text = " ".join((text or "").split())
    if len(text) <= n:
        return text
    cut = text[:n].rsplit(" ", 1)[0]
    return cut + "..."


def load_rows(path, keep_div01=False):
    with open(path, newline="", encoding="utf-8") as fh:
        raw = list(csv.DictReader(fh))
    rows = []
    for r in raw:
        full = (r.get("requirement") or r.get("description") or "").strip()
        desc = (r.get("description") or "").strip() or shorten(full)
        src = r.get("origin") or r.get("source") or ""
        rows.append({
            "sub_no": (r.get("sub_no") or "").strip(),
            "section": (r.get("section") or "").strip(),
            "section_title": (r.get("section_title") or "").strip(),
            "description": desc,
            "full": full,
            "types": (r.get("types") or "").strip(),
            "ref": " ".join(x for x in [
                (r.get("article") or "").strip(),
                f"p.{r['page']}" if r.get("page") else ""] if x),
            "origin": src,
            "lead": (r.get("lead_weeks") or "").strip(),
            "review": (r.get("review_weeks") or "").strip(),
            "notes": (r.get("notes") or "").strip(),
            "needs_review": truthy(r.get("needs_review")) or src.upper().startswith("REVIEW"),
            "review_note": (r.get("review_note") or "").strip(),
        })
    rows = [r for r in rows if r["section"] and r["description"]]
    if not keep_div01:  # Division 01 procedure text becomes project rules, not register rows
        rows = [r for r in rows if r["origin"].strip().lower() != "division 01 procedure"]
    return rows


def number_rows(rows, scheme):
    counters = defaultdict(int)
    seq = 0
    used = {r["sub_no"] for r in rows if r["sub_no"]}
    for r in rows:
        if r["sub_no"]:
            continue
        if scheme == "compact":  # 061000.01, the section digits then a two-digit sequence
            counters[r["section"]] += 1
            digits = r["section"].replace(" ", "")
            r["sub_no"] = f"{digits}.{counters[r['section']]:02d}"
            continue
        if scheme == "sequential":
            seq += 1
            r["sub_no"] = f"{seq:03d}"
        else:
            counters[r["section"]] += 1
            n = counters[r["section"]]
            cand = f"{r['section']}-{n}"
            while cand in used:
                n += 1
                cand = f"{r['section']}-{n}"
            used.add(cand)
            r["sub_no"] = cand


def load_install(path):
    out = {}
    if not path:
        return out
    with open(path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            d = datetime.strptime(r["install_date"].strip(), "%Y-%m-%d").date() if r.get("install_date") else None
            lead = float(r["lead_weeks"]) if r.get("lead_weeks") else None
            out[r["section"].strip()] = (d, lead)
    return out


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = f(True, "FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def build(rows, out, rules, install):
    wb = Workbook()

    # ---------------- Project Rules
    wr = wb.active
    wr.title = "Project Rules"
    wr["A1"] = "Project Rules"
    wr["A1"].font = f(True, size=14)
    wr["A2"] = "Yellow cells are inputs. The Register sheet reads the review and prep weeks from here."
    wr["A2"].font = f(italic=True, color="595959")
    items = [
        ("Project", rules.get("project", "")),
        ("Default review time (weeks)", rules.get("review_weeks", 3)),
        ("Default prep time (weeks)", rules.get("prep_weeks", 1)),
        ("Numbering format", rules.get("numbering_note", "Section number, dash, sequence (for example 23 73 13-1). Resubmittals add a letter suffix (23 73 13-1a).")),
        ("Copies and format", rules.get("copies_format", "")),
        ("Source of rules", rules.get("rules_source", "")),
    ]
    for i, (k, v) in enumerate(items, start=4):
        wr.cell(row=i, column=1, value=k).font = f(True)
        c = wr.cell(row=i, column=2, value=v)
        c.font = f(color="0000FF" if isinstance(v, (int, float)) else "000000")
        c.fill = INPUT_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BORDER
    wr["A11"] = "Project-wide submittal requirements"
    wr["A11"].font = f(True)
    for j, line in enumerate(rules.get("division01_notes", []), start=12):
        wr.cell(row=j, column=1, value=f"{j - 11}.").font = f()
        c = wr.cell(row=j, column=2, value=line)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wr.column_dimensions["A"].width = 34
    wr.column_dimensions["B"].width = 110
    REVIEW_REF = "'Project Rules'!$B$5"
    PREP_REF = "'Project Rules'!$B$6"

    # ---------------- Register
    ws = wb.create_sheet("Register", 0)
    heads = ["Sub #", "Rev", "Description", "Spec Section", "Section Title", "Type", "Spec Ref",
             "Install Date", "Lead Time (wk)", "Review (wk)", "Prep (wk)", "Required Approval",
             "Required Submit", "Actual Submit", "Approval Date", "Status", "Days to Due",
             "Origin", "Notes", "Spec Text"]
    H = {h: i + 1 for i, h in enumerate(heads)}
    L = {h: get_column_letter(i) for h, i in H.items()}
    ws["A1"] = (rules.get("project") or "Project") + " - Submittal Register"
    ws["A1"].font = f(True, size=14)
    ws["A2"] = ("Yellow cells are inputs. Required Approval = Install Date - Lead Time. "
                "Required Submit = Install Date - (Lead + Review + Prep). Orange rows need a human decision (see Review Queue).")
    ws["A2"].font = f(italic=True, color="595959")
    HR = 4
    for h, i in H.items():
        ws.cell(row=HR, column=i, value=h)
    style_header(ws, HR, len(heads))
    ws.row_dimensions[HR].height = 32

    r0 = HR + 1
    for n, r in enumerate(rows):
        x = r0 + n
        inst_date, inst_lead = install.get(r["section"], (None, None))
        lead = r["lead"] or (inst_lead if inst_lead is not None else "")
        try:
            lead = float(lead) if lead != "" else ""
        except ValueError:
            lead = ""
        vals = {
            "Sub #": r["sub_no"], "Rev": 0, "Description": r["description"],
            "Spec Section": r["section"], "Section Title": r["section_title"].title() if r["section_title"].isupper() else r["section_title"],
            "Type": r["types"], "Spec Ref": r["ref"],
            "Install Date": inst_date, "Lead Time (wk)": lead,
            "Review (wk)": float(r["review"]) if r["review"] else f"={REVIEW_REF}",
            "Prep (wk)": f"={PREP_REF}",
            "Required Approval": f'=IF(OR({L["Install Date"]}{x}="",{L["Lead Time (wk)"]}{x}=""),"",{L["Install Date"]}{x}-7*{L["Lead Time (wk)"]}{x})',
            "Required Submit": f'=IF(OR({L["Install Date"]}{x}="",{L["Lead Time (wk)"]}{x}=""),"",{L["Install Date"]}{x}-7*({L["Lead Time (wk)"]}{x}+{L["Review (wk)"]}{x}+{L["Prep (wk)"]}{x}))',
            "Actual Submit": None, "Approval Date": None, "Status": "Not Started",
            "Days to Due": f'=IF(OR({L["Required Submit"]}{x}="",{L["Status"]}{x}="Approved",{L["Status"]}{x}="Approved as Noted"),"",{L["Required Submit"]}{x}-TODAY())',
            "Origin": r["origin"],
            "Notes": r["notes"] or (r["review_note"] if r["needs_review"] else ""),
            "Spec Text": r["full"],
        }
        for h, v in vals.items():
            c = ws.cell(row=x, column=H[h], value=v)
            c.font = f()
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=h in ("Description", "Notes", "Section Title"))
            if h in ("Install Date", "Actual Submit", "Approval Date", "Required Approval", "Required Submit"):
                c.number_format = "mm/dd/yyyy"
            if h in ("Install Date", "Lead Time (wk)", "Actual Submit", "Approval Date", "Status", "Notes"):
                c.fill = INPUT_FILL
            if h in ("Lead Time (wk)",) and isinstance(v, float):
                c.font = f(color="0000FF")
            if h == "Review (wk)" and isinstance(v, str):
                c.font = f(color="008000")
            if h == "Prep (wk)":
                c.font = f(color="008000")
            if h == "Days to Due":
                c.number_format = "0;[Red]-0"
        if r["needs_review"]:
            for h in ("Sub #", "Description", "Spec Section", "Type", "Spec Ref", "Origin"):
                ws.cell(row=x, column=H[h]).fill = REVIEW_FILL
    last = r0 + len(rows) - 1

    widths = {"Sub #": 14, "Rev": 5, "Description": 58, "Spec Section": 12, "Section Title": 26, "Type": 11,
              "Spec Ref": 14, "Install Date": 12, "Lead Time (wk)": 10, "Review (wk)": 9, "Prep (wk)": 9,
              "Required Approval": 12, "Required Submit": 12, "Actual Submit": 12, "Approval Date": 12,
              "Status": 17, "Days to Due": 9, "Origin": 22, "Notes": 34, "Spec Text": 60}
    for h, w in widths.items():
        ws.column_dimensions[L[h]].width = w
    ws.freeze_panes = ws.cell(row=r0, column=H["Description"] + 1)
    ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(heads))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HR}:{HR}"
    ws.print_area = f"A1:{L['Status']}{last}"

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    dv.error, dv.errorTitle = "Pick a status from the list.", "Status"
    ws.add_data_validation(dv)
    dv.add(f'{L["Status"]}{r0}:{L["Status"]}{last}')

    stat = f'{L["Status"]}{r0}:{L["Status"]}{last}'
    green = PatternFill("solid", bgColor="E2EFDA", fgColor="E2EFDA")
    red = PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")
    blue = PatternFill("solid", bgColor="DDEBF7", fgColor="DDEBF7")
    ws.conditional_formatting.add(stat, FormulaRule(formula=[f'OR({L["Status"]}{r0}="Approved",{L["Status"]}{r0}="Approved as Noted")'], fill=green))
    ws.conditional_formatting.add(stat, FormulaRule(formula=[f'OR({L["Status"]}{r0}="Revise/Resubmit",{L["Status"]}{r0}="Rejected")'], fill=red))
    ws.conditional_formatting.add(stat, FormulaRule(formula=[f'{L["Status"]}{r0}="Submitted"'], fill=blue))
    due = f'{L["Days to Due"]}{r0}:{L["Days to Due"]}{last}'
    ws.conditional_formatting.add(due, CellIsRule(operator="lessThan", formula=["0"], fill=red))

    # ---------------- Review Queue
    wq = wb.create_sheet("Review Queue")
    wq["A1"] = "Review Queue"
    wq["A1"].font = f(True, size=14)
    wq["A2"] = "Items the extraction could not place with confidence. Confirm each, then edit or delete its row on the Register sheet."
    wq["A2"].font = f(italic=True, color="595959")
    qh = ["Sub #", "Spec Section", "Section Title", "Spec Ref", "Type", "Why it needs review", "Spec text", "Decision"]
    for i, h in enumerate(qh, start=1):
        wq.cell(row=4, column=i, value=h)
    style_header(wq, 4, len(qh))
    q = [r for r in rows if r["needs_review"]]
    for n, r in enumerate(q, start=5):
        why = r["review_note"] or ("Requirement found outside a submittals article" if "buried" in r["origin"].lower()
                                   else "Typed item in a section with no submittals article" if "no article" in r["origin"].lower()
                                   else "Flagged for review")
        for i, v in enumerate([r["sub_no"], r["section"], r["section_title"], r["ref"], r["types"], why, r["full"], ""], start=1):
            c = wq.cell(row=n, column=i, value=v)
            c.font = f()
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=i in (3, 6, 7))
            if i == 8:
                c.fill = INPUT_FILL
    for i, w in enumerate([14, 12, 26, 14, 10, 40, 90, 18], start=1):
        wq.column_dimensions[get_column_letter(i)].width = w
    wq.freeze_panes = "A5"
    if not q:
        wq["A5"] = "No items flagged."
        wq["A5"].font = f(italic=True)

    # ---------------- Summary by section
    wsu = wb.create_sheet("By Section")
    wsu["A1"] = "Submittals by Spec Section"
    wsu["A1"].font = f(True, size=14)
    for i, h in enumerate(["Spec Section", "Section Title", "Submittals", "Approved", "Open"], start=1):
        wsu.cell(row=3, column=i, value=h)
    style_header(wsu, 3, 5)
    secs = OrderedDict()
    for r in rows:
        secs.setdefault(r["section"], r["section_title"])
    reg = "Register"
    for n, (sec, title) in enumerate(secs.items(), start=4):
        wsu.cell(row=n, column=1, value=sec).font = f()
        wsu.cell(row=n, column=2, value=title.title() if title.isupper() else title).font = f()
        a = f"{reg}!${L['Spec Section']}${r0}:${L['Spec Section']}${last}"
        s = f"{reg}!${L['Status']}${r0}:${L['Status']}${last}"
        wsu.cell(row=n, column=3, value=f"=COUNTIF({a},A{n})").font = f()
        wsu.cell(row=n, column=4, value=f'=COUNTIFS({a},A{n},{s},"Approved")+COUNTIFS({a},A{n},{s},"Approved as Noted")').font = f()
        wsu.cell(row=n, column=5, value=f"=C{n}-D{n}").font = f()
        for c in range(1, 6):
            wsu.cell(row=n, column=c).border = BORDER
    tot = 4 + len(secs)
    wsu.cell(row=tot, column=1, value="Total").font = f(True)
    for col in (3, 4, 5):
        cl = get_column_letter(col)
        c = wsu.cell(row=tot, column=col, value=f"=SUM({cl}4:{cl}{tot - 1})")
        c.font = f(True)
        c.border = BORDER
    for i, w in enumerate([14, 44, 12, 12, 10], start=1):
        wsu.column_dimensions[get_column_letter(i)].width = w
    wsu.freeze_panes = "A4"

    # ---------------- Legend
    wl = wb.create_sheet("Legend")
    wl["A1"] = "Legend"
    wl["A1"].font = f(True, size=14)
    wl["A3"], wl["B3"] = "Type code", "Meaning"
    style_header(wl, 3, 2)
    for n, (k, v) in enumerate(TYPE_LEGEND, start=4):
        wl.cell(row=n, column=1, value=k).font = f(True)
        wl.cell(row=n, column=2, value=v).font = f()
    n = 4 + len(TYPE_LEGEND) + 1
    wl.cell(row=n, column=1, value="Cell colors").font = f(True)
    for k, (fill, text) in enumerate([(INPUT_FILL, "Input: edit these"), (REVIEW_FILL, "Needs a human decision"),
                                      (None, "Blue text: typed number. Green text: value linked from Project Rules. Black: formula or text.")], start=1):
        c = wl.cell(row=n + k, column=1, value="")
        if fill:
            c.fill = fill
        wl.cell(row=n + k, column=2, value=text).font = f()
    wl.column_dimensions["A"].width = 14
    wl.column_dimensions["B"].width = 80

    wb.save(out)
    return len(rows), len(q)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("csv")
    ap.add_argument("out")
    ap.add_argument("--rules")
    ap.add_argument("--project")
    ap.add_argument("--review-weeks", type=float)
    ap.add_argument("--prep-weeks", type=float)
    ap.add_argument("--numbering", choices=["section", "sequential", "compact"], default="section")
    ap.add_argument("--install-dates")
    ap.add_argument("--keep-div01", action="store_true",
                    help="keep rows tagged 'Division 01 procedure' (dropped by default)")
    a = ap.parse_args()

    rules = {}
    if a.rules:
        with open(a.rules, encoding="utf-8") as fh:
            rules = json.load(fh)
    if a.project:
        rules["project"] = a.project
    if a.review_weeks is not None:
        rules["review_weeks"] = a.review_weeks
    if a.prep_weeks is not None:
        rules["prep_weeks"] = a.prep_weeks
    rules.setdefault("review_weeks", 3)
    rules.setdefault("prep_weeks", 1)

    rows = load_rows(a.csv, a.keep_div01)
    if not rows:
        sys.exit("No usable rows in the CSV (need section and description or requirement).")
    number_rows(rows, a.numbering)
    n, q = build(rows, a.out, rules, load_install(a.install_dates))
    print(f"Wrote {a.out}: {n} submittals, {q} flagged for review")


if __name__ == "__main__":
    main()
