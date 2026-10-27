"""Google Sheets output: a JSON file of cell values ready to push through the Sheets connector.

Python cannot call the connector, so this adapter writes the values and the few settings that
need a separate call (status dropdown, date columns). Claude then creates the spreadsheet and
sends each tab's `values` with the USER_ENTERED option, so formulas and dates are parsed the
way they would be if someone typed them.

Layout matches the Excel register where it matters: the Project Rules tab keeps the review
weeks in B5 and the prep weeks in B6, and the Register tab calculates dates with formulas that
read them. The Register header is on row 1, so rows can be appended and filtered as a table.
"""
import json

from openpyxl.utils import get_column_letter

import columns as colmod
import register_model as model

REVIEW_REF = "'Project Rules'!$B$5"
PREP_REF = "'Project Rules'!$B$6"


def _rules_tab(rules):
    grid = [["Project Rules", ""], ["Yellow cells are inputs. The Register tab reads the review and prep weeks from here.", ""], ["", ""]]
    labels = [
        ("Project", rules.get("project", "")),
        ("Default review time (weeks)", rules.get("review_weeks", 3)),
        ("Default prep time (weeks)", rules.get("prep_weeks", 1)),
        ("Numbering format", rules.get("numbering_note", "")),
        ("Copies and format", rules.get("copies_format", "")),
        ("Source of rules", rules.get("rules_source", "")),
    ]
    grid += [[k, v] for k, v in labels]
    grid += [["", ""], ["Project-wide submittal requirements", ""]]
    grid += [[f"{i}.", line] for i, line in enumerate(rules.get("division01_notes", []), start=1)]
    return grid


def write(register, out, layout):
    rules, items, cols = register["rules"], register["items"], layout["columns"]
    L = {c["key"]: get_column_letter(i + 1) for i, c in enumerate(cols)}
    grid = [[c["header"] for c in cols]]
    for n, item in enumerate(items):
        x = n + 2  # data starts on row 2
        row = []
        for c in cols:
            k = c["key"]
            if k == "review":
                v = model.num(item["review"])
                v = v if v is not None else f"={REVIEW_REF}"
            elif k == "prep":
                v = f"={PREP_REF}"
            elif k == "required_approval":
                v = f'=IF(OR({L["install_date"]}{x}="",{L["lead"]}{x}=""),"",{L["install_date"]}{x}-7*{L["lead"]}{x})'
            elif k == "required_submit":
                v = (f'=IF(OR({L["install_date"]}{x}="",{L["lead"]}{x}=""),"",'
                     f'{L["install_date"]}{x}-7*({L["lead"]}{x}+{L["review"]}{x}+{L["prep"]}{x}))')
            elif k == "days_to_due":
                v = (f'=IF(OR({L["required_submit"]}{x}="",{L["status"]}{x}="Approved",'
                     f'{L["status"]}{x}="Approved as Noted"),"",{L["required_submit"]}{x}-TODAY())')
            elif k == "lead":
                n_lead = model.num(item["lead"])
                v = "" if n_lead is None else n_lead
            else:
                v = _plain_value(item, k)
            row.append(v)
        grid.append(row)
    last = len(items) + 1
    payload = {
        "note": "Create a spreadsheet, then send each tab's values with valueInputOption USER_ENTERED.",
        "tabs": {
            "Register": {"range": f"A1:{get_column_letter(len(cols))}{last}", "values": grid},
            "Project Rules": {"range": f"A1:B{len(_rules_tab(rules))}", "values": _rules_tab(rules)},
        },
        "date_columns": [L[k] for k in colmod.DATE_KEYS if k in L],
        "status_dropdown": {"column": L.get("status"), "first_row": 2, "last_row": last, "values": model.STATUSES},
        "frozen_rows": 1,
    }
    with open(out, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return len(items), sum(1 for i in items if i["needs_review"])


def _plain_value(item, key):
    """Plain value for one column key."""
    if key == "section_title":
        t = item["section_title"]
        return t.title() if t.isupper() else t
    if key == "notes":
        return item["notes"] or (item["review_note"] if item["needs_review"] else "")
    if key in item:
        return item[key]
    return item["extra"].get(key, "")
