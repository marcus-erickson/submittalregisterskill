"""Excel output and input for the register."""
from collections import OrderedDict
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import columns as colmod
import register_model as model

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INPUT_KEYS = ("install_date", "lead", "actual_submit", "approval_date", "status", "notes", "ball_in_court")
WRAP_KEYS = ("description", "notes", "section_title")
REVIEW_FILL_KEYS = ("sub_no", "description", "section", "types", "ref", "origin")
RULE_LABELS = ["Project", "Default review time (weeks)", "Default prep time (weeks)",
               "Numbering format", "Copies and format", "Source of rules"]


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = f(True, "FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _date(v):
    d = model.parse_date(v)
    return d if d else None


def _weeks(v):
    n = model.num(v)
    return n if n is not None else ""


# ------------------------------------------------------------------ write

def write(register, out, layout):
    rows = register["items"]
    rules = register["rules"]
    wb = Workbook()

    # ---------------- Project Rules
    wr = wb.active
    wr.title = "Project Rules"
    wr["A1"] = "Project Rules"
    wr["A1"].font = f(True, size=14)
    wr["A2"] = "Yellow cells are inputs. The Register sheet reads the review and prep weeks from here."
    wr["A2"].font = f(italic=True, color="595959")
    items = [
        ("Project", rules.get("project", "")),
        ("Default review time (weeks)", rules.get("review_weeks", 3)),
        ("Default prep time (weeks)", rules.get("prep_weeks", 1)),
        ("Numbering format", rules.get("numbering_note", "Section number, dash, sequence (for example 23 73 13-1). Resubmittals add a letter suffix (23 73 13-1a).")),
        ("Copies and format", rules.get("copies_format", "")),
        ("Source of rules", rules.get("rules_source", "")),
    ]
    for i, (k, v) in enumerate(items, start=4):
        wr.cell(row=i, column=1, value=k).font = f(True)
        c = wr.cell(row=i, column=2, value=v)
        c.font = f(color="0000FF" if isinstance(v, (int, float)) else "000000")
        c.fill = INPUT_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BORDER
    wr["A11"] = "Project-wide submittal requirements"
    wr["A11"].font = f(True)
    for j, line in enumerate(rules.get("division01_notes", []), start=12):
        wr.cell(row=j, column=1, value=f"{j - 11}.").font = f()
        c = wr.cell(row=j, column=2, value=line)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wr.column_dimensions["A"].width = 34
    wr.column_dimensions["B"].width = 110
    REVIEW_REF = "'Project Rules'!$B$5"
    PREP_REF = "'Project Rules'!$B$6"

    # ---------------- Register
    ws = wb.create_sheet("Register", 0)
    cols = layout["columns"]
    H = {c["key"]: i + 1 for i, c in enumerate(cols)}
    L = {k: get_column_letter(i) for k, i in H.items()}
    ws["A1"] = (rules.get("project") or "Project") + " - Submittal Register"
    ws["A1"].font = f(True, size=14)
    ws["A2"] = ("Yellow cells are inputs. Required Approval = Install Date - Lead Time. "
                "Required Submit = Install Date - (Lead + Review + Prep). Orange rows need a human decision (see Review Queue).")
    ws["A2"].font = f(italic=True, color="595959")
    HR = 4
    for c in cols:
        ws.cell(row=HR, column=H[c["key"]], value=c["header"])
    style_header(ws, HR, len(cols))
    ws.row_dimensions[HR].height = 32

    def has(*keys):
        return all(k in H for k in keys)

    r0 = HR + 1
    for n, r in enumerate(rows):
        x = r0 + n
        vals = {}
        for c in cols:
            k = c["key"]
            if k == "sub_no":
                v = r["sub_no"]
            elif k == "rev":
                v = r["rev"]
            elif k == "description":
                v = r["description"]
            elif k == "section":
                v = r["section"]
            elif k == "section_title":
                v = r["section_title"].title() if r["section_title"].isupper() else r["section_title"]
            elif k == "types":
                v = r["types"]
            elif k == "ref":
                v = r["ref"]
            elif k == "install_date":
                v = _date(r["install_date"])
            elif k == "lead":
                v = _weeks(r["lead"])
            elif k == "review":
                n_review = model.num(r["review"])
                v = n_review if n_review is not None else f"={REVIEW_REF}"
            elif k == "prep":
                v = f"={PREP_REF}"
            elif k == "required_approval":
                v = (f'=IF(OR({L["install_date"]}{x}="",{L["lead"]}{x}=""),"",'
                     f'{L["install_date"]}{x}-7*{L["lead"]}{x})')
            elif k == "required_submit":
                v = (f'=IF(OR({L["install_date"]}{x}="",{L["lead"]}{x}=""),"",'
                     f'{L["install_date"]}{x}-7*({L["lead"]}{x}+{L["review"]}{x}+{L["prep"]}{x}))')
            elif k == "actual_submit":
                v = _date(r["actual_submit"])
            elif k == "approval_date":
                v = _date(r["approval_date"])
            elif k == "status":
                v = r["status"] or "Not Started"
            elif k == "days_to_due":
                v = (f'=IF(OR({L["required_submit"]}{x}="",{L["status"]}{x}="Approved",'
                     f'{L["status"]}{x}="Approved as Noted"),"",{L["required_submit"]}{x}-TODAY())')
            elif k == "origin":
                v = r["origin"]
            elif k == "notes":
                v = r["notes"] or (r["review_note"] if r["needs_review"] else "")
            elif k == "full":
                v = r["full"]
            elif k == "ball_in_court":
                v = r["ball_in_court"]
            else:  # x_ custom column
                v = r["extra"].get(k, "")
            vals[k] = v
        for k, v in vals.items():
            c = ws.cell(row=x, column=H[k], value=v)
            c.font = f()
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=k in WRAP_KEYS)
            if k in colmod.DATE_KEYS:
                c.number_format = "mm/dd/yyyy"
            if k in INPUT_KEYS or k.startswith("x_"):
                c.fill = INPUT_FILL
            if k == "lead" and isinstance(v, float):
                c.font = f(color="0000FF")
            if k == "review" and isinstance(v, str):
                c.font = f(color="008000")
            if k == "prep":
                c.font = f(color="008000")
            if k == "days_to_due":
                c.number_format = "0;[Red]-0"
        if r["needs_review"]:
            for k in REVIEW_FILL_KEYS:
                if k in H:
                    ws.cell(row=x, column=H[k]).fill = REVIEW_FILL
    last = r0 + len(rows) - 1

    for c in cols:
        ws.column_dimensions[L[c["key"]]].width = c["width"]
    after = layout.get("freeze_after")
    ws.freeze_panes = ws.cell(row=r0, column=(H[after] + 1) if after in H else 1)
    ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(cols))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HR}:{HR}"
    through = layout.get("print_through")
    ws.print_area = f"A1:{L[through] if through in L else get_column_letter(len(cols))}{last}"

    red = PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")
    if "status" in H:
        sc = L["status"]
        dv = DataValidation(type="list", formula1='"' + ",".join(model.STATUSES) + '"', allow_blank=False)
        dv.error, dv.errorTitle = "Pick a status from the list.", "Status"
        ws.add_data_validation(dv)
        dv.add(f"{sc}{r0}:{sc}{last}")
        stat = f"{sc}{r0}:{sc}{last}"
        green = PatternFill("solid", bgColor="E2EFDA", fgColor="E2EFDA")
        blue = PatternFill("solid", bgColor="DDEBF7", fgColor="DDEBF7")
        ws.conditional_formatting.add(stat, FormulaRule(formula=[f'OR({sc}{r0}="Approved",{sc}{r0}="Approved as Noted")'], fill=green))
        ws.conditional_formatting.add(stat, FormulaRule(formula=[f'OR({sc}{r0}="Revise/Resubmit",{sc}{r0}="Rejected")'], fill=red))
        ws.conditional_formatting.add(stat, FormulaRule(formula=[f'{sc}{r0}="Submitted"'], fill=blue))
    if "days_to_due" in H:
        dc = L["days_to_due"]
        ws.conditional_formatting.add(f"{dc}{r0}:{dc}{last}", CellIsRule(operator="lessThan", formula=["0"], fill=red))

    # ---------------- Review Queue
    wq = wb.create_sheet("Review Queue")
    wq["A1"] = "Review Queue"
    wq["A1"].font = f(True, size=14)
    wq["A2"] = "Items the extraction could not place with confidence. Confirm each, then edit or delete its row on the Register sheet."
    wq["A2"].font = f(italic=True, color="595959")
    qh = ["Sub #", "Spec Section", "Section Title", "Spec Ref", "Type", "Why it needs review", "Spec text", "Decision"]
    for i, h in enumerate(qh, start=1):
        wq.cell(row=4, column=i, value=h)
    style_header(wq, 4, len(qh))
    q = [r for r in rows if r["needs_review"]]
    for n, r in enumerate(q, start=5):
        why = r["review_note"] or ("Requirement found outside a submittals article" if "buried" in r["origin"].lower()
                                   else "Typed item in a section with no submittals article" if "no article" in r["origin"].lower()
                                   else "Flagged for review")
        for i, v in enumerate([r["sub_no"], r["section"], r["section_title"], r["ref"], r["types"], why, r["full"], ""], start=1):
            c = wq.cell(row=n, column=i, value=v)
            c.font = f()
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=i in (3, 6, 7))
            if i == 8:
                c.fill = INPUT_FILL
    for i, w in enumerate([14, 12, 26, 14, 10, 40, 90, 18], start=1):
        wq.column_dimensions[get_column_letter(i)].width = w
    wq.freeze_panes = "A5"
    if not q:
        wq["A5"] = "No items flagged."
        wq["A5"].font = f(italic=True)

    # ---------------- Summary by section (needs the section and status columns)
    if has("section", "status"):
        wsu = wb.create_sheet("By Section")
        wsu["A1"] = "Submittals by Spec Section"
        wsu["A1"].font = f(True, size=14)
        for i, h in enumerate(["Spec Section", "Section Title", "Submittals", "Approved", "Open"], start=1):
            wsu.cell(row=3, column=i, value=h)
        style_header(wsu, 3, 5)
        secs = OrderedDict()
        for r in rows:
            secs.setdefault(r["section"], r["section_title"])
        reg = "Register"
        for n, (sec, title) in enumerate(secs.items(), start=4):
            wsu.cell(row=n, column=1, value=sec).font = f()
            wsu.cell(row=n, column=2, value=title.title() if title.isupper() else title).font = f()
            a = f"{reg}!${L['section']}${r0}:${L['section']}${last}"
            s = f"{reg}!${L['status']}${r0}:${L['status']}${last}"
            wsu.cell(row=n, column=3, value=f"=COUNTIF({a},A{n})").font = f()
            wsu.cell(row=n, column=4, value=f'=COUNTIFS({a},A{n},{s},"Approved")+COUNTIFS({a},A{n},{s},"Approved as Noted")').font = f()
            wsu.cell(row=n, column=5, value=f"=C{n}-D{n}").font = f()
            for c in range(1, 6):
                wsu.cell(row=n, column=c).border = BORDER
        tot = 4 + len(secs)
        wsu.cell(row=tot, column=1, value="Total").font = f(True)
        for col in (3, 4, 5):
            cl = get_column_letter(col)
            c = wsu.cell(row=tot, column=col, value=f"=SUM({cl}4:{cl}{tot - 1})")
            c.font = f(True)
            c.border = BORDER
        for i, w in enumerate([14, 44, 12, 12, 10], start=1):
            wsu.column_dimensions[get_column_letter(i)].width = w
        wsu.freeze_panes = "A4"

    # ---------------- Legend
    wl = wb.create_sheet("Legend")
    wl["A1"] = "Legend"
    wl["A1"].font = f(True, size=14)
    wl["A3"], wl["B3"] = "Type code", "Meaning"
    style_header(wl, 3, 2)
    for n, (k, v) in enumerate(model.TYPE_LEGEND, start=4):
        wl.cell(row=n, column=1, value=k).font = f(True)
        wl.cell(row=n, column=2, value=v).font = f()
    n = 4 + len(model.TYPE_LEGEND) + 1
    wl.cell(row=n, column=1, value="Cell colors").font = f(True)
    for k, (fill, text) in enumerate([(INPUT_FILL, "Input: edit these"), (REVIEW_FILL, "Needs a human decision"),
                                      (None, "Blue text: typed number. Green text: value linked from Project Rules. Black: formula or text.")], start=1):
        c = wl.cell(row=n + k, column=1, value="")
        if fill:
            c.fill = fill
        wl.cell(row=n + k, column=2, value=text).font = f()
    wl.column_dimensions["A"].width = 14
    wl.column_dimensions["B"].width = 80

    wb.save(out)
    return len(rows), len(q)


# ------------------------------------------------------------------ read

def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return model.iso(v)
    return str(v).strip()


def _find_header(ws, layout):
    best = (0, None, {})
    for row in range(1, min(ws.max_row, 20) + 1):
        headers = [_cell_text(c.value) for c in ws[row]]
        found, score = colmod.header_map(headers, layout)
        if score > best[0]:
            best = (score, row, found)
    score, row, found = best
    if row is None or not all(k in found.values() for k in colmod.CORE):
        raise ValueError("Could not find a header row with at least a section and a description column. "
                         "Pass a --columns layout whose headers match this workbook.")
    return row, found


def _read_rules(wb):
    rules = {}
    if "Project Rules" not in wb.sheetnames:
        return rules
    ws = wb["Project Rules"]
    keymap = {"Project": "project", "Default review time (weeks)": "review_weeks",
              "Default prep time (weeks)": "prep_weeks", "Numbering format": "numbering_note",
              "Copies and format": "copies_format", "Source of rules": "rules_source"}
    for row in range(4, 10):
        label = _cell_text(ws.cell(row=row, column=1).value)
        value = ws.cell(row=row, column=2).value
        if label in keymap and value not in (None, ""):
            rules[keymap[label]] = value
    notes = []
    for row in range(12, ws.max_row + 1):
        v = ws.cell(row=row, column=2).value
        if v not in (None, ""):
            notes.append(str(v))
    if notes:
        rules["division01_notes"] = notes
    return rules


def read(path, layout):
    """Load a register workbook into the model. Calculated columns are ignored."""
    wb = load_workbook(path)
    ws = wb["Register"] if "Register" in wb.sheetnames else wb.worksheets[0]
    header_row, found = _find_header(ws, layout)
    items = []
    for row in range(header_row + 1, ws.max_row + 1):
        vals, sub_cell = {}, None
        for pos, key in found.items():
            cell = ws.cell(row=row, column=pos + 1)
            vals[key] = cell.value
            if key == "sub_no":
                sub_cell = cell
        sec = _cell_text(vals.get("section"))
        desc = _cell_text(vals.get("description"))
        if not sec and not desc:
            continue
        item = model.blank_item(
            section=sec, description=desc,
            sub_no=_cell_text(vals.get("sub_no")),
            rev=model.to_int(vals.get("rev")),
            section_title=_cell_text(vals.get("section_title")),
            types=_cell_text(vals.get("types")),
            ref=_cell_text(vals.get("ref")),
            origin=_cell_text(vals.get("origin")),
            notes=_cell_text(vals.get("notes")),
            full=_cell_text(vals.get("full")),
            install_date=model.iso(vals.get("install_date")),
            actual_submit=model.iso(vals.get("actual_submit")),
            approval_date=model.iso(vals.get("approval_date")),
            status=_cell_text(vals.get("status")) or "Not Started",
            ball_in_court=_cell_text(vals.get("ball_in_court")),
        )
        # a formula in lead or review means "use the default"; only typed numbers are data
        for key in ("lead", "review"):
            n = model.num(vals.get(key))
            item[key] = n if n is not None else ""
        item["extra"] = {k: _cell_text(v) for k, v in vals.items() if k.startswith("x_")}
        # the orange fill on the Sub # cell is the live "needs a human decision" flag
        fill = sub_cell.fill if sub_cell is not None else None
        rgb = str(fill.fgColor.rgb) if fill is not None and fill.fill_type else ""
        item["needs_review"] = rgb.upper().endswith(REVIEW_FILL.fgColor.rgb[-6:].upper())
        if item["needs_review"]:
            item["review_note"] = item["notes"]
        items.append(item)
    return model.new_register(_read_rules(wb).get("project", ""), _read_rules(wb), items)
